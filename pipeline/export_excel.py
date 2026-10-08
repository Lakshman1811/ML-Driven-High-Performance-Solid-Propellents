from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

from .config import OUTPUT_DIR

LEFT_FILL = PatternFill("solid", fgColor="1F4E79")
RIGHT_FILL = PatternFill("solid", fgColor="C45911")
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)

DISPLAY_COLS = [
    "rank", "formula", "smiles", "Al_wt%", "EC_wt%", "HTPB_wt%", "AP_wt%",
    "Isp_s", "Tc_K", "Cstar_m_s", "O_over_C", "MW", "promising_Isp_gt_270",
]


def _write_table(ws, start_col, title, fill, df: pd.DataFrame, start_row=1):
    n = len(DISPLAY_COLS)
    ws.merge_cells(start_row=start_row, start_column=start_col, end_row=start_row, end_column=start_col + n - 1)
    cell = ws.cell(start_row, start_col, title)
    cell.fill = fill
    cell.font = HEADER_FONT
    cell.alignment = Alignment(horizontal="center")
    for j, name in enumerate(DISPLAY_COLS, start=start_col):
        c = ws.cell(start_row + 1, j, name)
        c.fill = fill
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    view = df[DISPLAY_COLS]
    for i, row in enumerate(view.itertuples(index=False), start=start_row + 2):
        for j, value in enumerate(row, start=start_col):
            cell = ws.cell(i, j, _excel_value(value))
            cell.border = THIN
            if isinstance(value, float):
                cell.number_format = "0.000"
    return start_col + n


def _excel_value(value):
    if isinstance(value, bool) or value is True or value is False:
        return bool(value)
    if hasattr(value, "item"):
        try:
            return value.item()
        except Exception:
            return value
    return value


def write_workbook(all_compounds, final7, metric_rows: list, loss_png, roc_png, paper_figs: dict, out_path: Path):
    wb = Workbook()

    # Sheet 1: All 1000+ Compounds (left) and Final 7 (right)
    ws = wb.active
    ws.title = "Screened_1000_and_Final7"
    _write_table(ws, 1, f"All {len(all_compounds):,} Screened Energetic Compounds (Ranked by GA Isp)", LEFT_FILL, all_compounds)
    _write_table(ws, 16, "Final 7 Promising ECs (Highest Isp; Paper Cutoff Isp > 270 s)", RIGHT_FILL, final7)
    ws.freeze_panes = "A3"
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["R"].width = 28

    # Sheet 2: Model Metrics Table (R2, MAE, RMSE across all models)
    ws_m = wb.create_sheet("Model_Metrics_R2_MAE_RMSE")
    ws_m["A1"] = "Comparative Model Performance (MLP vs Baselines: Ridge, AdaBoost, KNN)"
    ws_m["A1"].font = Font(bold=True, size=13, color="1F4E79")
    mdf = pd.DataFrame(metric_rows)
    headers = ["Model", "Target", "Train R²", "Train MAE", "Train RMSE", "Test R²", "Test MAE", "Test RMSE"]
    ws_m.append([])
    ws_m.append(headers)
    for col_idx in range(1, 9):
        cell = ws_m.cell(3, col_idx)
        cell.fill = LEFT_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    for r_idx, row in enumerate(mdf.itertuples(index=False), start=4):
        vals = [row.model.upper(), row.target, row.Train_R2, row.Train_MAE, row.Train_RMSE, row.Test_R2, row.Test_MAE, row.Test_RMSE]
        ws_m.append(vals)
        for c_idx in range(1, 9):
            c = ws_m.cell(r_idx, c_idx)
            c.border = THIN
            if c_idx > 2:
                c.number_format = "0.000"
                c.alignment = Alignment(horizontal="right")
            else:
                c.alignment = Alignment(horizontal="center")
    for col in ["A", "B", "C", "D", "E", "F", "G", "H"]:
        ws_m.column_dimensions[col].width = 15

    # Sheet 3: Loss Curves and ROC Curves
    ws_g = wb.create_sheet("Curves_and_ROC")
    ws_g["A1"] = "Train vs Test Loss (MLP Tc, Isp, C*)"
    ws_g["A1"].font = Font(bold=True, size=14, color="1F4E79")
    if Path(loss_png).is_file():
        img = XLImage(str(loss_png))
        img.width = 960
        img.height = 320
        ws_g.add_image(img, "A3")

    ws_g["A20"] = "ROC Curves (Isp, Tc, C*) with AUC"
    ws_g["A20"].font = Font(bold=True, size=14, color="1F4E79")
    if Path(roc_png).is_file():
        img2 = XLImage(str(roc_png))
        img2.width = 960
        img2.height = 320
        ws_g.add_image(img2, "A22")

    # Sheet 4: Paper Figures & SHAP
    ws_p = wb.create_sheet("Paper_Figures_and_SHAP")
    current_row = 1
    for title, fig_path in paper_figs.items():
        if fig_path and Path(fig_path).is_file():
            ws_p.cell(current_row, 1, title).font = Font(bold=True, size=13, color="1F4E79")
            img = XLImage(str(fig_path))
            img.width = 920
            img.height = int(920 * 0.55)
            ws_p.add_image(img, f"A{current_row + 2}")
            current_row += 24

    out_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(out_path)
    except PermissionError:
        alt_path = out_path.with_name(out_path.stem + "_latest.xlsx")
        print(f"[Notice] '{out_path.name}' is currently locked/opened. Saved workbook to '{alt_path.name}' instead.")
        wb.save(alt_path)
        return alt_path
    return out_path


